"""Testy doplňování hodnot a přebalení .docx."""

from __future__ import annotations

import io
import xml.dom.minidom
import zipfile

import pytest

from conftest import requires_sample
from dlg.docx_engine import extract_text, fill_docx, scan_docx
from dlg.docx_engine.parts import DocxError, DocxPackage
from fixtures import (
    PNG_1PX,
    document_xml,
    docx_via_python_docx,
    header_xml,
    make_docx,
    paragraph,
    part_bytes,
    part_text,
    run,
    sdt_date,
    table,
    text_paragraph,
    textbox,
)


def build(body: str, **kwargs: object) -> bytes:
    return make_docx(document_xml(body), **kwargs)  # type: ignore[arg-type]


def values_for(docx: bytes, mapping: dict[str, str]) -> dict[str, str]:
    """Namapuje hodnoty podle ``raw`` textu placeholderu na jeho id."""

    scan = scan_docx(docx)
    out: dict[str, str] = {}
    for placeholder in scan.placeholders:
        if placeholder.raw in mapping:
            out[placeholder.id] = mapping[placeholder.raw]
    assert len(out) == len(mapping), "některý placeholder se nenašel"
    return out


def assert_valid_xml(docx: bytes, part: str = "word/document.xml") -> str:
    text = part_text(docx, part)
    xml.dom.minidom.parseString(text)
    return text


# ---------------------------------------------------------------------------
# základní doplnění
# ---------------------------------------------------------------------------

def test_doplneni_jednoduche_hodnoty() -> None:
    docx = build(text_paragraph("Vážený [Jméno],"))
    out, report = fill_docx(docx, values_for(docx, {"[Jméno]": "Jan Novák"}))
    assert extract_text(out) == "Vážený Jan Novák,"
    assert len(report.filled) == 1
    assert report.unfilled == []
    assert report.warnings == []
    assert_valid_xml(out)


def test_hodnota_se_xml_escapuje() -> None:
    docx = build(text_paragraph("[X]"))
    out, _ = fill_docx(docx, values_for(docx, {"[X]": 'a & b < c > d "e"'}))
    assert "&amp; b &lt; c &gt;" in part_text(out)
    assert extract_text(out) == 'a & b < c > d "e"'


def test_hodnota_s_diakritikou_a_nedelitelnou_mezerou() -> None:
    docx = build(text_paragraph("[X]"))
    value = "Příliš\xa0žluťoučký kůň ● úpěl"
    out, _ = fill_docx(docx, values_for(docx, {"[X]": value}))
    assert extract_text(out) == value


def test_hodnota_rozdelena_do_vice_behu() -> None:
    body = paragraph(run("Pan ["), run("Jan Novák", highlight="yellow"), run("] z Prahy"))
    docx = build(body)
    out, _ = fill_docx(docx, values_for(docx, {"[Jan Novák]": "Petr Dvořák"}))
    text = assert_valid_xml(out)
    assert extract_text(out) == "Pan Petr Dvořák z Prahy"
    # hodnota jde do prvního w:t rozsahu, ostatní se vyprázdní
    assert "<w:t>Pan Petr Dvořák</w:t>" in text
    assert text.count("<w:t") == 3


def test_doplneni_zachova_okolni_text_a_prida_xml_space() -> None:
    body = paragraph(run("["), run("A"), run("] zbytek věty."))
    docx = build(body)
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "X"}))
    text = assert_valid_xml(out)
    assert '<w:t xml:space="preserve"> zbytek věty.</w:t>' in text
    assert extract_text(out) == "X zbytek věty."


def test_hodnota_koncici_mezerou_dostane_xml_space() -> None:
    docx = build(text_paragraph("[A]"))
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "  odsazeno  "}))
    text = assert_valid_xml(out)
    assert 'xml:space="preserve"' in text
    assert extract_text(out) == "  odsazeno  "


def test_novy_radek_se_prevede_na_zalomeni() -> None:
    docx = build(text_paragraph("[A]"))
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "první\ndruhý\r\ntřetí"}))
    text = assert_valid_xml(out)
    assert text.count("<w:br/>") == 2
    assert '</w:t><w:br/><w:t xml:space="preserve">druhý' in text
    assert extract_text(out) == "první\ndruhý\ntřetí"


def test_ridici_znaky_se_zahodi() -> None:
    docx = build(text_paragraph("[A]"))
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "a\x07b\x00c"}))
    assert extract_text(out) == "abc"
    assert_valid_xml(out)


# ---------------------------------------------------------------------------
# zvýraznění
# ---------------------------------------------------------------------------

def test_zvyrazneni_se_odstrani() -> None:
    body = paragraph(run("["), run("Jan", highlight="yellow"), run("]"))
    docx = build(body)
    out, _ = fill_docx(docx, values_for(docx, {"[Jan]": "Petr"}))
    assert "w:highlight" not in part_text(out)


def test_zvyrazneni_lze_ponechat() -> None:
    body = paragraph(run("["), run("Jan", highlight="yellow"), run("]"))
    docx = build(body)
    out, _ = fill_docx(docx, values_for(docx, {"[Jan]": "Petr"}), clear_highlight=False)
    assert 'w:highlight w:val="yellow"' in part_text(out)


def test_zvyrazneni_v_ppr_se_take_odstrani() -> None:
    body = paragraph(
        run("[A]", highlight="yellow"),
        ppr_extra='<w:rPr><w:highlight w:val="green"/></w:rPr>',
    )
    docx = build(body)
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "X"}))
    assert "w:highlight" not in part_text(out)


def test_zvyrazneni_jinde_zustane() -> None:
    body = text_paragraph("[A]") + paragraph(run("jiný", highlight="cyan"))
    docx = build(body)
    scan = scan_docx(docx)
    target = [p for p in scan.placeholders if p.raw == "[A]"][0]
    out, _ = fill_docx(docx, {target.id: "X"})
    assert part_text(out).count("w:highlight") == 1


# ---------------------------------------------------------------------------
# sdt
# ---------------------------------------------------------------------------

def test_sdt_se_pri_zapisu_rozbali() -> None:
    docx = build(sdt_date("6. září 2026"))
    out, report = fill_docx(docx, values_for(docx, {"6. září 2026": "9. září 2026"}))
    text = assert_valid_xml(out)
    assert "<w:sdt>" not in text and "</w:sdt>" not in text
    assert "w:sdtContent" not in text and "w:sdtPr" not in text
    assert "<w:p" in text
    assert extract_text(out) == "9. září 2026"
    assert len(report.filled) == 1


def test_nevyplneny_sdt_zustane() -> None:
    docx = build(sdt_date("6. září 2026"))
    out, report = fill_docx(docx, {})
    assert "<w:sdt>" in part_text(out)
    assert len(report.unfilled) == 1


# ---------------------------------------------------------------------------
# nevyplněné placeholdery
# ---------------------------------------------------------------------------

def test_keep_unfilled_true_nechava_beze_zmeny() -> None:
    docx = build(text_paragraph("[A] a [B]"))
    values = values_for(docx, {"[A]": "X"})
    out, report = fill_docx(docx, values, keep_unfilled=True)
    assert extract_text(out) == "X a [B]"
    assert len(report.filled) == 1 and len(report.unfilled) == 1


def test_keep_unfilled_false_placeholder_smaze() -> None:
    docx = build(text_paragraph("[A] a [B]"))
    values = values_for(docx, {"[A]": "X"})
    out, report = fill_docx(docx, values, keep_unfilled=False)
    assert extract_text(out) == "X a "
    assert len(report.unfilled) == 1
    assert_valid_xml(out)


def test_prazdna_hodnota_se_bere_jako_nevyplneno() -> None:
    docx = build(text_paragraph("[A]"))
    values = values_for(docx, {"[A]": "   "})
    out, report = fill_docx(docx, values)
    assert extract_text(out) == "[A]"
    assert report.filled == [] and len(report.unfilled) == 1


def test_neznamy_placeholder_da_varovani() -> None:
    docx = build(text_paragraph("[A]"))
    _, report = fill_docx(docx, {"ph_neexistuje": "X"})
    assert any("neznámý placeholder" in w for w in report.warnings)


# ---------------------------------------------------------------------------
# vypuštění odstavců
# ---------------------------------------------------------------------------

def test_vypusteni_odstavce() -> None:
    docx = build(text_paragraph("první") + text_paragraph("druhý") + text_paragraph("třetí"))
    scan = scan_docx(docx)
    target = [p for p in scan.paragraphs if p.text == "druhý"][0]
    out, report = fill_docx(docx, {}, drop_paragraphs=[target.id])
    assert extract_text(out) == "první\ntřetí"
    assert report.dropped_paragraphs == [target.id]
    assert_valid_xml(out)


def test_vypusteni_posledniho_odstavce_v_bunce_jen_vyprazdni() -> None:
    body = table([[[text_paragraph("jediný v buňce")], [text_paragraph("vedle")]]])
    docx = make_docx(document_xml(body))
    scan = scan_docx(docx)
    target = [p for p in scan.paragraphs if p.text == "jediný v buňce"][0]
    out, report = fill_docx(docx, {}, drop_paragraphs=[target.id])
    text = assert_valid_xml(out)
    assert text.count("<w:tc>") == 2
    assert "jediný v buňce" not in text
    assert any("buňce tabulky" in w for w in report.warnings)
    assert target.id in report.dropped_paragraphs
    # v buňce musí zůstat aspoň jeden odstavec
    assert "<w:tc><w:tcPr/><w:p" in text


def test_vypusteni_jednoho_ze_dvou_odstavcu_v_bunce_smaze_cely() -> None:
    body = table([[[text_paragraph("první"), text_paragraph("druhý")]]])
    docx = make_docx(document_xml(body))
    scan = scan_docx(docx)
    target = [p for p in scan.paragraphs if p.text == "první"][0]
    out, report = fill_docx(docx, {}, drop_paragraphs=[target.id])
    text = assert_valid_xml(out)
    assert "první" not in text and "druhý" in text
    assert report.warnings == []


def test_vypusteni_odstavce_s_placeholderem() -> None:
    docx = build(text_paragraph("[A]") + text_paragraph("zůstává"))
    scan = scan_docx(docx)
    target = [p for p in scan.paragraphs if p.text == "[A]"][0]
    placeholder = scan.placeholders[0]
    out, report = fill_docx(docx, {placeholder.id: "X"}, drop_paragraphs=[target.id])
    assert extract_text(out) == "zůstává"
    assert report.filled == []
    assert any("vypuštěn" in w for w in report.warnings)


def test_neznamy_odstavec_da_varovani() -> None:
    docx = build(text_paragraph("A"))
    _, report = fill_docx(docx, {}, drop_paragraphs=["pg_neexistuje"])
    assert any("pg_neexistuje" in w for w in report.warnings)
    assert report.dropped_paragraphs == []


# ---------------------------------------------------------------------------
# více částí, ZIP
# ---------------------------------------------------------------------------

def test_doplneni_v_zahlavi() -> None:
    docx = make_docx(
        document_xml(text_paragraph("tělo [A]")),
        {"word/header1.xml": header_xml(text_paragraph("záhlaví [B]"))},
    )
    out, report = fill_docx(docx, values_for(docx, {"[A]": "X", "[B]": "Y"}))
    assert "záhlaví Y" in part_text(out, "word/header1.xml")
    assert "tělo X" in part_text(out)
    assert len(report.filled) == 2


def test_zip_zustane_kompletni_a_ve_stejnem_poradi() -> None:
    docx = build(text_paragraph("[A]"))
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "X"}))
    with zipfile.ZipFile(io.BytesIO(docx)) as before, zipfile.ZipFile(io.BytesIO(out)) as after:
        assert before.namelist() == after.namelist()
        for info in before.infolist():
            if info.filename == "word/document.xml":
                continue
            assert before.read(info.filename) == after.read(info.filename)
        assert after.testzip() is None
        image = [i for i in after.infolist() if i.filename.endswith(".png")][0]
        assert image.compress_type == zipfile.ZIP_STORED
    assert part_bytes(out, "word/media/image1.png") == PNG_1PX


def test_slozkove_polozky_prezijou() -> None:
    docx = build(text_paragraph("[A]"), with_directory_entry=True)
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "X"}))
    with zipfile.ZipFile(io.BytesIO(out)) as archive:
        assert "word/" in archive.namelist()
        assert archive.getinfo("word/").is_dir()


def test_glossary_se_nemeni() -> None:
    docx = build(text_paragraph("[Jméno]"))
    out, _ = fill_docx(docx, values_for(docx, {"[Jméno]": "Petr"}))
    assert part_bytes(out, "word/glossary/document.xml") == part_bytes(
        docx, "word/glossary/document.xml"
    )


def test_bez_zmen_zustanou_casti_bajtove_stejne() -> None:
    docx = build(text_paragraph("[A]"))
    out, report = fill_docx(docx, {})
    assert part_bytes(out, "word/document.xml") == part_bytes(docx, "word/document.xml")
    assert len(report.unfilled) == 1


def test_nevalidni_soubor() -> None:
    with pytest.raises(DocxError):
        fill_docx(b"tohle rozhodne neni zip", {})
    with pytest.raises(DocxError):
        scan_docx(make_docx(document_xml(""))[:20])
    prazdny_zip = io.BytesIO()
    with zipfile.ZipFile(prazdny_zip, "w") as archive:
        archive.writestr("neco.txt", b"x")
    with pytest.raises(DocxError, match="word/document.xml"):
        DocxPackage.open(prazdny_zip.getvalue())


def test_chybejici_soubor(tmp_path) -> None:
    with pytest.raises(DocxError, match="neexistuje"):
        scan_docx(tmp_path / "nic.docx")


# ---------------------------------------------------------------------------
# textová pole, opakované skenování
# ---------------------------------------------------------------------------

def test_doplneni_v_textovem_poli() -> None:
    docx = make_docx(document_xml(textbox(text_paragraph("[A]"))))
    scan = scan_docx(docx)
    values = {p.id: "X" for p in scan.placeholders}
    out, report = fill_docx(docx, values)
    assert len(report.filled) == 2  # mc:Choice i mc:Fallback
    assert "[A]" not in part_text(out)
    assert_valid_xml(out)


def test_vysledek_jde_znovu_naskenovat() -> None:
    docx = build(text_paragraph("[A] a [B]"))
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "X", "[B]": "Y"}))
    assert scan_docx(out).placeholders == []


def test_opakovane_vyplneni_je_stabilni() -> None:
    docx = build(text_paragraph("[A]"))
    values = values_for(docx, {"[A]": "X"})
    first, _ = fill_docx(docx, values)
    second, _ = fill_docx(docx, values)
    assert part_bytes(first, "word/document.xml") == part_bytes(second, "word/document.xml")


def test_vice_placeholderu_v_jednom_bezi() -> None:
    docx = build(text_paragraph("[A] a [B] a [C]"))
    out, report = fill_docx(docx, values_for(docx, {"[A]": "1", "[B]": "2", "[C]": "3"}))
    assert extract_text(out) == "1 a 2 a 3"
    assert len(report.filled) == 3
    assert_valid_xml(out)


def test_fill_prijme_cestu(tmp_path) -> None:
    docx = build(text_paragraph("[A]"))
    path = tmp_path / "s.docx"
    path.write_bytes(docx)
    out, _ = fill_docx(path, values_for(docx, {"[A]": "X"}))
    assert extract_text(out) == "X"


# ---------------------------------------------------------------------------
# balíček od python-docx
# ---------------------------------------------------------------------------

def test_vyplneny_balicek_otevre_python_docx() -> None:
    import docx as python_docx

    source = docx_via_python_docx()
    values = values_for(
        source,
        {
            "6. září 2026": "9. září 2026",
            "[Jan Novák]": "Petr Dvořák",
            "{{ domena }}": "lego-shop.cz",
            "[Petr Novák]": "Karel Vomáčka",
            "[Česká republika / Slovenská republika]": "Slovenská republika",
        },
    )
    out, report = fill_docx(source, values)
    assert len(report.filled) == 5
    document = python_docx.Document(io.BytesIO(out))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Petr Dvořák" in text
    assert "lego-shop.cz" in text
    cells = [c.text for row in document.tables[0].rows for c in row.cells]
    assert "Držitel: Karel Vomáčka" in cells
    assert "w:highlight" not in part_text(out)


# ---------------------------------------------------------------------------
# extract_text
# ---------------------------------------------------------------------------

def test_extract_text_vraci_odstavce_hlavniho_dokumentu() -> None:
    docx = make_docx(
        document_xml(text_paragraph("první") + text_paragraph("") + text_paragraph("třetí")),
        {"word/header1.xml": header_xml(text_paragraph("záhlaví"))},
    )
    assert extract_text(docx) == "první\n\ntřetí"


# ---------------------------------------------------------------------------
# reálná šablona
# ---------------------------------------------------------------------------

@requires_sample
def test_realna_sablona_end_to_end(sample_docx_path) -> None:
    scan = scan_docx(sample_docx_path)
    hodnoty = {
        "[Jan Novák]": "Petr Dvořák & syn",
        "[Nová 1]": "Krátká 12",
        "[110 00]": "602 00",
        "[Česká republika / Slovenská republika]": "Slovenská republika",
        "[●]": "lego-shop.cz",
    }
    values: dict[str, str] = {}
    for placeholder in scan.placeholders:
        if placeholder.raw in hodnoty:
            values[placeholder.id] = hodnoty[placeholder.raw]
        elif placeholder.kind == "sdt":
            values[placeholder.id] = "9. září 2026"
        elif placeholder.raw.startswith("[zrušili registraci"):
            values[placeholder.id] = placeholder.options[1]
    assert len(values) == 7

    out, report = fill_docx(sample_docx_path, values)
    assert len(report.filled) == 7
    assert report.unfilled == []
    assert report.warnings == []

    # výsledek je platné XML ve všech skenovaných částech
    package = DocxPackage.open(out)
    for name in package.scan_parts():
        xml.dom.minidom.parseString(package.read(name))

    text = part_text(out)
    assert "w:highlight" not in text
    assert "<w:sdt>" not in text

    # jde znovu naskenovat a nic nezbylo
    assert scan_docx(out).placeholders == []

    plain = extract_text(out)
    for value in ("Petr Dvořák & syn", "Krátká 12", "602 00", "Slovenská republika",
                  "lego-shop.cz", "9. září 2026"):
        assert value in plain
    for original in ("Jan Novák", "Nová 1", "110 00", "[●]",
                     "Česká republika / Slovenská republika", "zrušili registraci",
                     "neděle 6. září 2026"):
        assert original not in plain

    # kromě word/document.xml se nic nezměnilo
    with zipfile.ZipFile(io.BytesIO(sample_docx_path.read_bytes())) as before:
        with zipfile.ZipFile(io.BytesIO(out)) as after:
            assert before.namelist() == after.namelist()
            changed = [
                name
                for name in before.namelist()
                if before.read(name) != after.read(name)
            ]
    assert changed == ["word/document.xml"]


@requires_sample
def test_realna_sablona_otevre_python_docx(sample_docx_bytes) -> None:
    import docx as python_docx

    scan = scan_docx(sample_docx_bytes)
    values = {p.id: (p.options[0] if p.options else "DOPLNĚNO") for p in scan.placeholders}
    out, _ = fill_docx(sample_docx_bytes, values)
    document = python_docx.Document(io.BytesIO(out))
    assert any("DOPLNĚNO" in p.text for p in document.paragraphs)


@requires_sample
def test_realna_sablona_bez_hodnot_zustane_bajtove_stejna(sample_docx_bytes) -> None:
    out, report = fill_docx(sample_docx_bytes, {})
    assert len(report.unfilled) == 7
    assert part_bytes(out, "word/document.xml") == part_bytes(
        sample_docx_bytes, "word/document.xml"
    )


@requires_sample
def test_realna_sablona_smaze_nevyplnene(sample_docx_bytes) -> None:
    out, report = fill_docx(sample_docx_bytes, {}, keep_unfilled=False)
    plain = extract_text(out)
    assert len(report.unfilled) == 7
    for original in ("[Jan Novák]", "[110 00]", "[●]"):
        assert original not in plain
    assert scan_docx(out).placeholders == []
    xml.dom.minidom.parseString(part_text(out))


def test_vypusteni_odstavce_s_textovym_polem() -> None:
    body = textbox(text_paragraph("[Uvnitř]")) + text_paragraph("zůstává")
    docx = make_docx(document_xml(body))
    scan = scan_docx(docx)
    outer = [p for p in scan.paragraphs if p.text == ""][0]
    values = {p.id: "X" for p in scan.placeholders}
    out, report = fill_docx(docx, values, drop_paragraphs=[outer.id])
    text = assert_valid_xml(out)
    assert "[Uvnitř]" not in text and "X" not in text
    assert extract_text(out) == "zůstává"
    assert report.filled == []


def test_extract_text_prevede_zalomeni_a_tabulator() -> None:
    body = paragraph(
        "<w:r><w:t>a</w:t><w:br/><w:t>b</w:t><w:tab/><w:t>c</w:t></w:r>"
    )
    docx = make_docx(document_xml(body))
    assert extract_text(docx) == "a\nb\tc"


def test_extract_text_ignoruje_tabulatory_v_ppr() -> None:
    body = paragraph(
        run("a"),
        ppr_extra='<w:tabs><w:tab w:val="left" w:pos="1134"/></w:tabs>',
    )
    docx = make_docx(document_xml(body))
    assert extract_text(docx) == "a"


def test_vicaradkova_hodnota_v_nahledu() -> None:
    docx = build(text_paragraph("[A]"))
    out, _ = fill_docx(docx, values_for(docx, {"[A]": "první\ndruhý"}))
    assert extract_text(out) == "první\ndruhý"
